import json
import io
import openpyxl
import pytest
from esh_savings.pipeline.ingestion import load_session
from esh_savings.pipeline import Pipeline
from esh_savings.models.config import AnalysisConfig
from esh_savings.models.se_inputs import SEProvidedInputs
from esh_savings.pipeline.standards.jurisdictions.us import USAdapter
from esh_savings.reporting.json import result_to_dict, result_to_json
from esh_savings.reporting.excel import result_to_excel


@pytest.fixture
def computed_result(sample_hdf5_path):
    session = load_session(sample_hdf5_path)
    pipeline = Pipeline(
        config=AnalysisConfig(),
        se_inputs=SEProvidedInputs(operator_count=5),
        adapter=USAdapter(),
    )
    return pipeline.run(session)


def test_result_to_dict_is_json_serializable(computed_result):
    d = result_to_dict(computed_result)
    assert isinstance(d, dict)
    assert "esh_risk_score_manual" in d
    raw = json.dumps(d)
    restored = json.loads(raw)
    assert restored["esh_risk_score_manual"] == pytest.approx(d["esh_risk_score_manual"])


def test_result_to_json_string(computed_result):
    s = result_to_json(computed_result)
    assert isinstance(s, str)
    obj = json.loads(s)
    assert obj["annual_esh_savings_usd"] > 0


def test_excel_returns_bytes(computed_result):
    wb_bytes = result_to_excel(computed_result)
    assert isinstance(wb_bytes, bytes)
    assert len(wb_bytes) > 0


def test_excel_has_three_sheets(computed_result):
    wb_bytes = result_to_excel(computed_result)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    assert "Summary" in wb.sheetnames
    assert "Exposure" in wb.sheetnames
    assert "Assumptions" in wb.sheetnames


def test_excel_exposure_sheet_has_a8_value(computed_result):
    wb_bytes = result_to_excel(computed_result)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["Exposure"]
    # Row 4 col C should be the A(8) value (a float)
    a8_cell = ws["C4"].value
    assert isinstance(a8_cell, float)


def test_excel_assumptions_has_shift_duration(computed_result):
    wb_bytes = result_to_excel(computed_result)
    wb = openpyxl.load_workbook(io.BytesIO(wb_bytes))
    ws = wb["Assumptions"]
    # Row 2 col A should label "Shift duration (s)"
    assert "shift" in str(ws["A2"].value).lower()
    assert ws["B2"].value == 28_800.0
